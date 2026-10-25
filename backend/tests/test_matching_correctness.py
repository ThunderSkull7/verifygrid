"""Unit tests for Batch 1: Reconciliation Correctness, Fuzzy Matching Edge Cases & Lineage Regression."""

from __future__ import annotations

import csv
import io
import openpyxl
import pytest

from app.schemas.result import (
    DiscrepancyCategory,
    MatchStatus,
    ReasonCode,
)
from app.services.export_service import (
    generate_csv_report,
    generate_xlsx_report,
    sanitize_formula_value,
)
from app.services.verification_engine import (
    check_initials_match,
    evaluate_name_similarity,
    evaluate_status_rule,
    normalize_name_tokens,
    verify_datasets,
)


# ==============================================================================
# 1. Name Variation Handling: Initials & Token Reordering
# ==============================================================================

def test_normalize_name_tokens():
    assert normalize_name_tokens("Doe, John A.") == ["doe", "john", "a"]
    assert normalize_name_tokens("  Mary - Jane   Watson ") == ["mary", "jane", "watson"]
    assert normalize_name_tokens(None) == []
    assert normalize_name_tokens("") == []


def test_check_initials_match():
    # Matching initials
    assert check_initials_match(["j", "doe"], ["john", "doe"]) is True
    assert check_initials_match(["john", "doe"], ["j", "doe"]) is True
    assert check_initials_match(["a", "j", "smith"], ["alice", "jane", "smith"]) is True

    # Non-initials or mismatched
    assert check_initials_match(["bob", "doe"], ["john", "doe"]) is False
    assert check_initials_match(["b", "doe"], ["john", "doe"]) is False
    assert check_initials_match([], []) is False


def test_evaluate_name_similarity():
    # Exact case-insensitive
    is_cand, score, reason = evaluate_name_similarity("John Doe", "john doe")
    assert is_cand is True
    assert score == 100.0

    # Token reordering: "John Doe" vs "Doe, John"
    is_cand, score, reason = evaluate_name_similarity("John Doe", "Doe, John")
    assert is_cand is True
    assert score == 100.0
    assert reason == "token_reorder"

    # Initials variation: "J. Doe" vs "John Doe"
    is_cand, score, reason = evaluate_name_similarity("J. Doe", "John Doe")
    assert is_cand is True
    assert score >= 90.0
    assert reason == "initials_variation"

    # Reversed initials: "Doe, J." vs "John Doe"
    is_cand, score, reason = evaluate_name_similarity("Doe, J.", "John Doe")
    assert is_cand is True
    assert score >= 90.0

    # Complete mismatch
    is_cand, score, reason = evaluate_name_similarity("Alice Smith", "Zachary Taylor")
    assert is_cand is False
    assert score < 50.0


def test_name_variation_enters_needs_review():
    """Initials and token-order variations must enter NEEDS_REVIEW rather than hard match/mismatch."""
    source = [
        (1, {"id": "101", "name": "Doe, John", "role": "Engineer"}),
        (2, {"id": "102", "name": "J. Smith", "role": "Analyst"}),
    ]
    target = [
        (1, {"id": "101", "name": "John Doe", "role": "Engineer"}),
        (2, {"id": "102", "name": "Jane Smith", "role": "Analyst"}),
    ]

    summary, results, discrepancies, warnings, errors = verify_datasets(
        source_records=source,
        target_records=target,
        key_columns=["id"],
    )

    # Both records matched on ID, but have name variations
    assert summary.matched_records == 2
    assert summary.identical_matches == 0
    assert summary.needs_review_records == 2
    assert summary.field_mismatches == 0
    assert summary.total_discrepancies == 2

    # Check results match status & reason code
    item1 = [r for r in results if r.key_values["id"] == "101"][0]
    assert item1.match_status == MatchStatus.NEEDS_REVIEW
    assert item1.discrepancy_category == DiscrepancyCategory.NEEDS_REVIEW
    assert item1.reason_code == ReasonCode.FUZZY_NAME_REVIEW
    assert item1.confidence_score is not None and item1.confidence_score >= 0.9
    assert item1.source_row_index == 1
    assert item1.target_row_index == 1

    item2 = [r for r in results if r.key_values["id"] == "102"][0]
    assert item2.match_status == MatchStatus.NEEDS_REVIEW
    assert item2.reason_code == ReasonCode.FUZZY_NAME_REVIEW
    assert item2.source_row_index == 2
    assert item2.target_row_index == 2


# ==============================================================================
# 2. Conflicting Identifiers Never Overridden by High Fuzzy Score
# ==============================================================================

def test_conflicting_ids_with_similar_names_not_auto_matched():
    """Ensure conflicting primary keys/IDs are never overridden simply because name similarity is high."""
    source = [
        (1, {"emp_id": "EMP001", "full_name": "Alexander Hamilton", "dept": "Treasury"}),
    ]
    target = [
        (1, {"emp_id": "EMP999", "full_name": "Alexander Hamilton", "dept": "Treasury"}),
    ]

    summary, results, discrepancies, warnings, errors = verify_datasets(
        source_records=source,
        target_records=target,
        key_columns=["emp_id"],
        enable_fuzzy_matching=True,
        fuzzy_threshold=80.0,
    )

    # Because emp_id is conflicting (EMP001 != EMP999), it MUST NOT be auto-matched to IDENTICAL
    assert summary.identical_matches == 0
    assert summary.needs_review_records == 1
    assert summary.total_discrepancies == 1

    cand = discrepancies[0]
    assert cand.match_status == MatchStatus.NEEDS_REVIEW
    assert cand.discrepancy_category == DiscrepancyCategory.NEEDS_REVIEW
    assert cand.reason_code == ReasonCode.ID_CONFLICT
    assert cand.source_row_index == 1
    assert cand.target_row_index == 1
    assert "EMP001" in cand.message
    assert "EMP999" in cand.message

    # Check field difference contains explicit ID conflict
    diff_reasons = [d.reason_code for d in cand.field_differences]
    assert ReasonCode.ID_CONFLICT in diff_reasons


def test_conflicting_id_in_matched_record_flags_id_conflict():
    """If records match on another key, conflicting secondary ID columns must flag ID_CONFLICT."""
    source = [
        (1, {"email": "john@example.com", "national_id": "111-22-3333", "name": "John Doe"}),
    ]
    target = [
        (1, {"email": "john@example.com", "national_id": "999-88-7777", "name": "John Doe"}),
    ]

    summary, results, discrepancies, warnings, errors = verify_datasets(
        source_records=source,
        target_records=target,
        key_columns=["email"],
    )

    assert summary.identical_matches == 0
    assert summary.needs_review_records == 1
    disc = discrepancies[0]
    assert disc.match_status == MatchStatus.NEEDS_REVIEW
    assert disc.reason_code == ReasonCode.ID_CONFLICT


# ==============================================================================
# 3. State-Aware Verification: Active vs. Suspended Status Rules
# ==============================================================================

def test_evaluate_status_rule_scenarios():
    # Same status
    is_perm, msg = evaluate_status_rule("Active", "Active")
    assert is_perm is True
    assert msg is None

    # Mismatch without rules
    is_perm, msg = evaluate_status_rule("Active", "Suspended", status_rules=None)
    assert is_perm is False
    assert "Status mismatch" in msg

    # Mismatch with permit_suspended flag
    is_perm, msg = evaluate_status_rule("Active", "Suspended", status_rules={"permit_suspended": True})
    assert is_perm is True

    # Transition rule permitting active -> suspended
    rules = {"allowed_transitions": {"active": ["active", "suspended"]}}
    is_perm, msg = evaluate_status_rule("active", "suspended", status_rules=rules)
    assert is_perm is True

    # Transition rule prohibiting active -> terminated
    is_perm, msg = evaluate_status_rule("active", "terminated", status_rules=rules)
    assert is_perm is False
    assert "Status rule violation" in msg


def test_status_mismatch_blocks_verification_without_rule():
    """An enrollment with non-matching status must be flagged as STATUS_MISMATCH and cannot pass as verified."""
    source = [
        (1, {"student_id": "S100", "status": "active", "gpa": 3.8}),
    ]
    target = [
        (1, {"student_id": "S100", "status": "suspended", "gpa": 3.8}),
    ]

    # Run without status rule permitting suspended
    summary, results, discrepancies, warnings, errors = verify_datasets(
        source_records=source,
        target_records=target,
        key_columns=["student_id"],
    )

    assert summary.identical_matches == 0
    assert summary.field_mismatches == 1
    assert summary.total_discrepancies == 1

    disc = discrepancies[0]
    assert disc.match_status == MatchStatus.FIELD_MISMATCH
    assert disc.reason_code == ReasonCode.STATUS_MISMATCH
    diff = disc.field_differences[0]
    assert diff.difference_type == "status_mismatch"
    assert diff.reason_code == ReasonCode.STATUS_MISMATCH


def test_status_permitted_by_rule_passes_as_identical():
    """When status rules permit the transition, the record can pass as identical."""
    source = [
        (1, {"student_id": "S100", "status": "active", "gpa": 3.8}),
    ]
    target = [
        (1, {"student_id": "S100", "status": "suspended", "gpa": 3.8}),
    ]

    status_rules = {
        "allowed_transitions": {
            "active": ["suspended"],
        }
    }

    summary, results, discrepancies, warnings, errors = verify_datasets(
        source_records=source,
        target_records=target,
        key_columns=["student_id"],
        status_rules=status_rules,
    )

    # Status transition is permitted, so records match identically!
    assert summary.identical_matches == 1
    assert summary.field_mismatches == 0
    assert summary.total_discrepancies == 0
    assert results[0].match_status == MatchStatus.IDENTICAL
    assert results[0].reason_code == ReasonCode.EXACT_MATCH


# ==============================================================================
# 4. Formula Injection Sanitization (Core Security)
# ==============================================================================

def test_sanitize_formula_value():
    # Spreadsheet formula triggers
    assert sanitize_formula_value("=1+1") == "'=1+1"
    assert sanitize_formula_value("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert sanitize_formula_value("@SUM(A1:A10)") == "'@SUM(A1:A10)"
    assert sanitize_formula_value("\t=2+2") == "'\t=2+2"
    assert sanitize_formula_value("\r=cmd") == "'\r=cmd"

    # Non-numeric minus string
    assert sanitize_formula_value("-cmd") == "'-cmd"

    # Legitimate numbers should NOT be prepended with quote
    assert sanitize_formula_value(100) == 100
    assert sanitize_formula_value(-45.5) == -45.5
    assert sanitize_formula_value("-100.50") == "-100.50"
    assert sanitize_formula_value("+250") == "+250"

    # Regular strings
    assert sanitize_formula_value("Alice Smith") == "Alice Smith"
    assert sanitize_formula_value(None) == ""


def test_csv_and_xlsx_export_neutralizes_formula_injection():
    """Ensure exports neutralize formula injection payloads in CSV and XLSX."""
    from unittest.mock import MagicMock
    import uuid

    mock_run = MagicMock()
    mock_run.id = uuid.uuid4()
    mock_run.project_id = uuid.uuid4()
    mock_run.status = "completed"
    mock_run.key_columns = ["id"]
    mock_run.started_at = None
    mock_run.completed_at = None
    mock_run.summary = {"total_source_records": 1, "total_target_records": 1}

    mock_disc = MagicMock()
    mock_disc.discrepancy_category = "=HYPERLINK(\"http://evil.com\")"
    mock_disc.match_status = "+cmd|' /C calc'!A0"
    mock_disc.source_row_number = 1
    mock_disc.target_row_number = 1
    mock_disc.source_row_numbers = None
    mock_disc.target_row_numbers = None
    mock_disc.key_values = {"id": "101"}
    mock_disc.field_differences = [
        {
            "field_name": "comment",
            "source_value": "-exec(calc.exe)",
            "target_value": "=cmd",
            "difference_type": "value_mismatch",
        }
    ]
    mock_disc.message = "@SUM(A1:A10)"

    # 1. Test CSV Export
    csv_bytes = generate_csv_report(mock_run, [mock_disc])
    csv_text = csv_bytes.decode("utf-8")
    assert "'=HYPERLINK" in csv_text
    assert "'+cmd" in csv_text
    assert "'@SUM" in csv_text
    assert "'-exec" in csv_text
    assert "'=cmd" in csv_text

    # 2. Test XLSX Export
    xlsx_bytes = generate_xlsx_report(mock_run, [mock_disc])
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Discrepancies"]
    cell_cat = ws.cell(row=2, column=1).value
    cell_status = ws.cell(row=2, column=2).value
    cell_msg = ws.cell(row=2, column=7).value

    assert str(cell_cat).startswith("'=")
    assert str(cell_status).startswith("'+")
    assert str(cell_msg).startswith("'@")


# ==============================================================================
# 5. Deterministic Row Lineage & Reason Codes
# ==============================================================================

def test_deterministic_row_lineage_and_amount_tolerance():
    source = [
        (1, {"id": "TX100", "amount": 100.50, "note": "payment"}),
    ]
    target = [
        (1, {"id": "TX100", "amount": 150.00, "note": "payment"}),
    ]

    summary, results, discrepancies, warnings, errors = verify_datasets(
        source_records=source,
        target_records=target,
        key_columns=["id"],
        numeric_tolerance=1.0,
        source_file_id="src-file-uuid",
        target_file_id="tgt-file-uuid",
    )

    assert summary.field_mismatches == 1
    assert len(discrepancies) == 1
    item = discrepancies[0]

    assert item.source_row_index == 1
    assert item.target_row_index == 1
    assert item.source_file_id == "src-file-uuid"
    assert item.target_file_id == "tgt-file-uuid"
    assert item.reason_code == ReasonCode.AMOUNT_TOLERANCE_EXCEEDED
    assert item.lineage["key_columns"] == ["id"]
    assert item.lineage["source_row_index"] == 1
    assert item.lineage["target_row_index"] == 1
